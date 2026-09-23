"""Offline smoke tests for user-visible report flows."""

import asyncio
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from bot.helpers import build_excel_from_context, send_long_text
from bot.helpers import parse_comparison_input, parse_period_input
from reports.generator import generate_comparison_periods_report
from reports.generator import generate_report


def test_period_parser_rejects_invalid_and_future_dates():
    today = date(2026, 9, 24)
    assert parse_period_input("01.09-07.09", today) == (
        date(2026, 9, 1), date(2026, 9, 7)
    )

    for value in ("31.02", "25.09", "07.09-01.09"):
        try:
            parse_period_input(value, today)
        except ValueError:
            pass
        else:
            raise AssertionError(f"invalid period accepted: {value}")


def test_comparison_parser_preserves_both_periods():
    result = parse_comparison_input(
        "01.09-07.09 vs 08.09-14.09", date(2026, 9, 24)
    )
    assert result == (
        date(2026, 9, 1), date(2026, 9, 7),
        date(2026, 9, 8), date(2026, 9, 14),
    )


def test_excel_dialogue_is_wrapped_and_readable():
    context = {
        "account": "@demo — Demo",
        "account_username": "demo",
        "period": "1–7 сентября",
        "stats": {},
        "content": {},
        "stories": {},
        "stories_list": [],
        "daily_stats": [],
        "publications": [],
        "ai_analysis": "Короткий анализ",
    }
    dialogue = [
        {"role": "user", "content": "Какой формат лучше?"},
        {"role": "assistant", "content": "Первая строка\nВторая строка"},
    ]

    workbook_bytes, filename = build_excel_from_context(context, dialogue)
    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook["Диалог с AI"]

    assert filename.endswith(".xlsx")
    assert sheet["B6"].alignment.wrap_text is True
    assert sheet["B6"].value == "Первая строка\nВторая строка"
    assert sheet.row_dimensions[6].height >= 24


def test_excel_comparison_contains_daily_and_format_sections():
    period = {"name": "1–7 сентября", "stats": {"reach_total": 700, "views_total": 1400, "accounts_engaged_total": 70}, "content": {}, "stories": {}}
    context = {
        "account": "@demo — Demo", "account_username": "demo", "period": "comparison",
        "period1": period, "period2": period,
        "period_days": {"available1": 7, "available2": 14},
        "comparison_formats": {"period1": {"REELS": {"posts": 2, "engagement_rate": 5.0}}, "period2": {}},
        "comparison_top_posts": {"period1": [{"caption": "Top", "value": 20}], "period2": []},
    }
    workbook_bytes, _ = build_excel_from_context(context, [])
    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook["Сравнение"]
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
    assert "Средние значения за день" in values
    assert "Форматы контента" in values
    assert "Лучшие публикации" in values


def test_send_long_text_never_sends_empty_chunks():
    class FakeMessage:
        def __init__(self):
            self.sent = []

        async def answer(self, text, **kwargs):
            self.sent.append(text)

    message = FakeMessage()
    asyncio.run(send_long_text(message, "x" * 4100))

    assert len(message.sent) == 2
    assert all(chunk.strip() for chunk in message.sent)
    assert all(len(chunk) <= 4096 for chunk in message.sent)


def test_comparison_rejects_overlapping_periods_before_api_call():
    account = SimpleNamespace(id=1, username="demo")

    async def run():
        await generate_comparison_periods_report(
            account,
            date(2026, 9, 1), date(2026, 9, 10),
            date(2026, 9, 10), date(2026, 9, 20),
        )

    try:
        asyncio.run(run())
    except ValueError as error:
        assert "пересекаться" in str(error)
    else:
        raise AssertionError("overlapping periods were accepted")


def test_full_report_pipeline_builds_context_from_mocked_data():
    account = SimpleNamespace(id=1, username="demo", name="Demo")
    stats = [SimpleNamespace(
        date=date(2026, 9, 1), followers=100, follower_count=3,
        reach=500, views=900, accounts_engaged=40, is_partial=False,
    )]
    post = SimpleNamespace(
        media_type="REELS", caption="A useful post", permalink="https://example.test/p",
        timestamp=datetime(2026, 9, 1, 12, 0),
        likes=20, comments=3, saved=4, shares=2, reach=500,
        total_interactions=29, views=800,
    )
    story = SimpleNamespace(
        timestamp=datetime(2026, 9, 1, 12, 0),
        media_type="IMAGE", permalink="", views=100, reach=80, replies=2,
        shares=1, total_interactions=3, profile_activity=1, follows=1,
        tap_forward=10, tap_back=2, tap_exit=3, swipe_forward=1, is_active=False,
    )

    async def run():
        with patch("reports.generator._fetch_and_save_data", new=AsyncMock(return_value={"api_delay_dates": [], "partial": False})), \
             patch("reports.generator._refresh_stories_if_recent", new=AsyncMock()), \
             patch("reports.generator.get_analyzer", return_value=None), \
             patch("reports.generator.crud.get_daily_stats", new=AsyncMock(return_value=stats)), \
             patch("reports.generator.crud.get_posts", new=AsyncMock(return_value=[post])), \
             patch("reports.generator.crud.get_stories", new=AsyncMock(return_value=[story])), \
             patch("reports.generator.create_followers_chart", return_value=b"followers"), \
             patch("reports.generator.create_metrics_chart", return_value=b"reach"), \
             patch("reports.generator.create_stories_chart", return_value=b"stories"):
            return await generate_report(account, custom_date_from=date(2026, 9, 1), custom_date_to=date(2026, 9, 1))

    result = asyncio.run(run())
    context = result["context"]
    assert result["text"]
    assert context["content"]["total_posts"] == 1
    assert context["top_posts"][0]["media_type"] == "REELS"
    assert context["format_performance"]["REELS"]["posts"] == 1
    assert context["data_quality"]["complete"] is True
    assert set(result["charts"]) == {"followers", "reach", "stories"}

