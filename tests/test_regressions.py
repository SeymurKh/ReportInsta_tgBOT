import io

from openpyxl import load_workbook

from bot.helpers import build_excel_from_context, parse_period_input


def test_parse_period_rejects_future_date():
    from datetime import date

    try:
        parse_period_input("22.09", date(2026, 9, 21))
    except ValueError:
        return
    raise AssertionError("future dates must be rejected")


def test_comparison_excel_contains_comparison_sheet():
    context = {
        "account_username": "test",
        "account": "@test — Test",
        "period": "1–7 сентября vs 8–14 сентября",
        "type": "comparison",
        "period1": {
            "name": "1–7 сентября",
            "stats": {"followers_growth": 10, "reach_total": 100},
            "content": {"total_posts": 2, "engagement_rate": 3.0},
            "stories": {"total_stories": 1, "total_views": 50},
        },
        "period2": {
            "name": "8–14 сентября",
            "stats": {"followers_growth": 20, "reach_total": 200},
            "content": {"total_posts": 4, "engagement_rate": 4.0},
            "stories": {"total_stories": 2, "total_views": 100},
        },
    }

    data, filename = build_excel_from_context(context, None)
    workbook = load_workbook(filename=io.BytesIO(data), read_only=True)

    assert filename.endswith(".xlsx")
    assert "Сравнение" in workbook.sheetnames
    comparison_sheet = workbook["Сравнение"]
    values = [cell.value for row in comparison_sheet.iter_rows() for cell in row]
    assert 10 in values
    assert 20 in values
