"""Polished Excel report generator using openpyxl."""

import io

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "17365D"
BLUE = "2196F3"
LIGHT_BLUE = "EAF3FF"
PALE_BLUE = "F5F9FF"
TEXT = "1F2937"
MUTED = "6B7280"
WHITE = "FFFFFF"
GRID = "D9E2F3"
THIN_BORDER = Border(
    left=Side(style="thin", color=GRID), right=Side(style="thin", color=GRID),
    top=Side(style="thin", color=GRID), bottom=Side(style="thin", color=GRID),
)


def _fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _hs(cell):
    cell.font = Font(bold=True, size=10, color=WHITE)
    cell.fill = _fill(BLUE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def _ss(cell):
    cell.font = Font(bold=True, size=11, color=NAVY)
    cell.fill = _fill(LIGHT_BLUE)
    cell.alignment = Alignment(vertical="center")
    cell.border = THIN_BORDER


def _title(ws, title, subtitle, last_col):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=18, color=WHITE)
    title_cell.fill = _fill(NAVY)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    subtitle_cell = ws.cell(row=2, column=1, value=subtitle)
    subtitle_cell.font = Font(size=10, color=MUTED, italic=True)
    subtitle_cell.fill = _fill(PALE_BLUE)
    subtitle_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 22


def _card(ws, label_cell, value_cell, label, value):
    for cell in (label_cell, value_cell):
        cell.fill = _fill(PALE_BLUE)
        cell.border = THIN_BORDER
    label_cell.value = label
    label_cell.font = Font(size=9, color=MUTED, bold=True)
    label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    value_cell.value = value
    value_cell.font = Font(size=16, color=NAVY, bold=True)
    value_cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_body(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.font = Font(size=10, color=TEXT)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (1, 3))
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = _fill(PALE_BLUE)


def _add_table(ws, ref, name):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)


def _set_numeric_format(ws, columns, min_row, max_row, fmt="#,##0"):
    for column in columns:
        for row in range(min_row, max_row + 1):
            ws.cell(row=row, column=column).number_format = fmt


def generate_excel_report(
    account_username, account_name, period_str,
    stats_summary, content_summary, daily_stats,
    publications, ai_analysis, dialogue_history=None,
    stories_summary=None, stories_list=None, comparison_periods=None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws.sheet_properties.tabColor = NAVY
    _title(ws, f"Отчёт: @{account_username}", f"{account_name}  •  Период: {period_str}", 4)
    ws.freeze_panes = "A5"
    for col, width in {"A": 29, "B": 17, "C": 29, "D": 17}.items():
        ws.column_dimensions[col].width = width

    ws.merge_cells("A4:D4")
    ws["A4"] = "Ключевые показатели"
    _ss(ws["A4"])
    for cell in ws["A4:D4"][0]:
        cell.fill = _fill(LIGHT_BLUE)
        cell.border = THIN_BORDER
    cards = [
        ("A5", "B5", "Подписчики", stats_summary.get("followers_end", "—")),
        ("C5", "D5", "Прирост", stats_summary.get("followers_growth", "—")),
        ("A7", "B7", "Охват", stats_summary.get("reach_total", "—")),
        ("C7", "D7", "Просмотры", stats_summary.get("views_total", "—")),
        ("A9", "B9", "Публикации", content_summary.get("total_posts", 0)),
        ("C9", "D9", "ER", content_summary.get("engagement_rate", 0)),
    ]
    for label_ref, value_ref, label, value in cards:
        _card(ws, ws[label_ref], ws[value_ref], label, value)
    for row in (5, 7, 9):
        ws.row_dimensions[row].height = 24
    ws["D9"].number_format = '0.0"%"'

    row = 12
    sections = [
        ("Аккаунт", [("Имя", account_name), ("Username", f"@{account_username}")]),
        ("Статистика", [
            ("Подписчики", stats_summary.get("followers_end", "—")),
            ("Прирост", stats_summary.get("followers_growth", "—")),
            ("Прирост %", stats_summary.get("followers_growth_pct", 0)),
            ("Охват", stats_summary.get("reach_total", "—")),
            ("Просмотры", stats_summary.get("views_total", "—")),
            ("Вовлечено", stats_summary.get("accounts_engaged_total", "—")),
            ("Ср. охват/день", stats_summary.get("reach_avg_daily", "—")),
        ]),
        ("Контент", [
            ("Публикаций", content_summary.get("total_posts", 0)), ("Reels", content_summary.get("total_reels", 0)),
            ("Видео", content_summary.get("total_videos", 0)), ("Фото", content_summary.get("total_images", 0)),
            ("Карусели", content_summary.get("total_carousels", 0)), ("Лайки", content_summary.get("total_likes", 0)),
            ("Комментарии", content_summary.get("total_comments", 0)), ("Сохранения", content_summary.get("total_saves", 0)),
            ("Репосты", content_summary.get("total_shares", 0)), ("Ср. лайки", content_summary.get("avg_likes", 0)),
            ("Ср. охват", content_summary.get("avg_reach", 0)), ("ER %", content_summary.get("engagement_rate", 0)),
        ]),
    ]
    for title, items in sections:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.cell(row=row, column=1, value=title)
        _ss(ws.cell(row=row, column=1))
        for col in (1, 2):
            ws.cell(row=row, column=col).border = THIN_BORDER
        for label, value in items:
            row += 1
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=value)
            _style_body(ws, row, row, 1, 2)
            if "%" in label:
                ws.cell(row=row, column=2).number_format = '0.0"%"'
            elif isinstance(value, (int, float)):
                ws.cell(row=row, column=2).number_format = "#,##0.0" if isinstance(value, float) else "#,##0"
        row += 2

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.cell(row=row, column=1, value="AI-анализ")
    _ss(ws.cell(row=row, column=1))
    for col in range(1, 5):
        ws.cell(row=row, column=col).border = THIN_BORDER
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 5, end_column=4)
    ai_cell = ws.cell(row=row, column=1, value=ai_analysis or "Нет данных")
    ai_cell.font = Font(size=10, color=TEXT)
    ai_cell.fill = _fill(PALE_BLUE)
    ai_cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cells in ws.iter_rows(min_row=row, max_row=row + 5, min_col=1, max_col=4):
        for cell in cells:
            cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 105

    # Daily statistics.
    ws2 = wb.create_sheet("По дням")
    ws2.sheet_properties.tabColor = BLUE
    _title(ws2, "Динамика по дням", f"@{account_username}  •  {period_str}", 5)
    for col, header in enumerate(["Дата", "Охват", "Прирост", "Просмотры", "Вовлечено"], 1):
        _hs(ws2.cell(row=4, column=col, value=header))
        ws2.column_dimensions[get_column_letter(col)].width = 18
    for i, data in enumerate(daily_stats, 5):
        for col, value in enumerate([data.get("date", ""), data.get("reach", 0), data.get("followers", 0), data.get("views", 0), data.get("accounts_engaged", 0)], 1):
            ws2.cell(row=i, column=col, value=value)
    if daily_stats:
        last_row = 4 + len(daily_stats)
        _style_body(ws2, 5, last_row, 1, 5)
        _set_numeric_format(ws2, range(2, 6), 5, last_row)
        _add_table(ws2, f"A4:E{last_row}", "DailyStats")
        ws2.conditional_formatting.add(f"C5:C{last_row}", ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    ws2.freeze_panes = "A5"

    # Publications.
    ws3 = wb.create_sheet("Публикации")
    ws3.sheet_properties.tabColor = "70AD47"
    _title(ws3, "Публикации", f"@{account_username}  •  {period_str}", 8)
    headers3 = ["Дата", "Тип", "Описание", "Лайки", "Комментарии", "Сохран.", "Репосты", "Охват"]
    for col, header in enumerate(headers3, 1):
        _hs(ws3.cell(row=4, column=col, value=header))
    for col, width in enumerate([14, 14, 52, 12, 14, 12, 12, 14], 1):
        ws3.column_dimensions[get_column_letter(col)].width = width
    type_names = {"IMAGE": "Фото", "VIDEO": "Видео", "CAROUSEL_ALBUM": "Карусель", "REELS": "Reels"}
    row = 5
    for day in publications:
        for post in day.get("posts", []):
            values = [day.get("date", ""), type_names.get(post.get("type", ""), post.get("type", "")), post.get("caption", "")[:160], post.get("likes", 0), post.get("comments", 0), post.get("saved", 0), post.get("shares", 0), post.get("reach", 0)]
            for col, value in enumerate(values, 1):
                ws3.cell(row=row, column=col, value=value)
            row += 1
    if row > 5:
        _style_body(ws3, 5, row - 1, 1, 8)
        _set_numeric_format(ws3, range(4, 9), 5, row - 1)
        _add_table(ws3, f"A4:H{row - 1}", "Publications")
    ws3.freeze_panes = "A5"

    # Stories.
    if stories_summary:
        ws_s = wb.create_sheet("Сторис")
        ws_s.sheet_properties.tabColor = "ED7D31"
        _title(ws_s, "Сторис", f"@{account_username}  •  {period_str}", 12)
        ws_s.column_dimensions["A"].width = 27
        ws_s.column_dimensions["B"].width = 16
        ws_s.merge_cells("A4:B4")
        ws_s["A4"] = "Сводка по сторис"
        _ss(ws_s["A4"])
        for col in (1, 2):
            ws_s.cell(row=4, column=col).border = THIN_BORDER
        story_metrics = [("Всего сторис", "total_stories"), ("Просмотры (итого)", "total_views"), ("Просмотры (средние)", "avg_views"), ("Охват (итого)", "total_reach"), ("Охват (средний)", "avg_reach"), ("Ответы", "total_replies"), ("Репосты", "total_shares"), ("Переходы в профиль", "total_profile_activity"), ("Подписки со сторис", "total_follows"), ("Выходы, %", "exit_rate"), ("Тапы вперёд", "tap_forward_total"), ("Тапы назад", "tap_back_total")]
        for i, (label, key) in enumerate(story_metrics, 5):
            ws_s.cell(row=i, column=1, value=label)
            ws_s.cell(row=i, column=2, value=stories_summary.get(key, 0))
            _style_body(ws_s, i, i, 1, 2)
            ws_s.cell(row=i, column=2).number_format = '0.0"%"' if " %" in label else "#,##0"
        header_row = 19
        headers_s = ["Дата", "Тип", "Просмотры", "Охват", "Ответы", "Репосты", "Профиль", "Подписки", "Вперёд", "Назад", "Выходы", "Свайпы"]
        for col, header in enumerate(headers_s, 1):
            _hs(ws_s.cell(row=header_row, column=col, value=header))
            ws_s.column_dimensions[get_column_letter(col)].width = max(ws_s.column_dimensions[get_column_letter(col)].width or 0, 13)
        story_type_names = {"IMAGE": "Фото", "VIDEO": "Видео"}
        for i, story in enumerate(stories_list or [], header_row + 1):
            values = [story.get("date", ""), story_type_names.get(story.get("media_type", ""), story.get("media_type", "")), story.get("views", 0), story.get("reach", 0), story.get("replies", 0), story.get("shares", 0), story.get("profile_activity", 0), story.get("follows", 0), story.get("tap_forward", 0), story.get("tap_back", 0), story.get("tap_exit", 0), story.get("swipe_forward", 0)]
            for col, value in enumerate(values, 1):
                ws_s.cell(row=i, column=col, value=value)
        if stories_list:
            last_story_row = header_row + len(stories_list)
            _style_body(ws_s, header_row + 1, last_story_row, 1, 12)
            _set_numeric_format(ws_s, range(3, 13), header_row + 1, last_story_row)
            _add_table(ws_s, f"A{header_row}:L{last_story_row}", "Stories")
        ws_s.freeze_panes = f"A{header_row + 1}"

    # AI dialogue.
    if dialogue_history:
        ws4 = wb.create_sheet("Диалог с AI")
        ws4.sheet_properties.tabColor = "A5A5A5"
        _title(ws4, "Диалог с AI", f"@{account_username}  •  {period_str}", 2)
        ws4.column_dimensions["A"].width = 16
        ws4.column_dimensions["B"].width = 72
        _hs(ws4.cell(row=4, column=1, value="Роль"))
        _hs(ws4.cell(row=4, column=2, value="Сообщение"))
        for i, message in enumerate(dialogue_history, 5):
            ws4.cell(row=i, column=1, value="Вопрос" if message.get("role") == "user" else "Ответ")
            content = str(message.get("content", ""))
            ws4.cell(row=i, column=2, value=content)
            # Approximate the displayed height so long AI answers remain readable.
            visual_lines = max(1, (len(content) // 95) + content.count("\n") + 1)
            ws4.row_dimensions[i].height = min(180, max(24, visual_lines * 15))
        _style_body(ws4, 5, 4 + len(dialogue_history), 1, 2)
        for i in range(5, 5 + len(dialogue_history)):
            ws4.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws4.freeze_panes = "A5"

    # Comparison.
    if comparison_periods:
        ws_cmp = wb.create_sheet("Сравнение")
        ws_cmp.sheet_properties.tabColor = "8064A2"
        _title(ws_cmp, "Сравнение периодов", f"@{account_username}", 3)
        for col, header in enumerate(["Показатель", "Период 1", "Период 2"], 1):
            _hs(ws_cmp.cell(row=4, column=col, value=header))
        p1, p2 = comparison_periods
        metric_groups = [("Статистика", [("Прирост подписчиков", "followers_growth"), ("Охват", "reach_total"), ("Просмотры", "views_total"), ("Вовлечено", "accounts_engaged_total")], "stats"), ("Контент", [("Публикации", "total_posts"), ("Лайки", "total_likes"), ("Комментарии", "total_comments"), ("ER %", "engagement_rate")], "content"), ("Stories", [("Количество", "total_stories"), ("Просмотры", "total_views"), ("Ответы", "total_replies")], "stories")]
        row = 5
        for group, metrics, key in metric_groups:
            ws_cmp.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            ws_cmp.cell(row=row, column=1, value=group)
            _ss(ws_cmp.cell(row=row, column=1))
            for col in range(1, 4):
                ws_cmp.cell(row=row, column=col).border = THIN_BORDER
            row += 1
            d1, d2 = p1.get(key, {}), p2.get(key, {})
            for label, metric in metrics:
                ws_cmp.cell(row=row, column=1, value=label)
                ws_cmp.cell(row=row, column=2, value=d1.get(metric, 0))
                ws_cmp.cell(row=row, column=3, value=d2.get(metric, 0))
                _style_body(ws_cmp, row, row, 1, 3)
                if metric == "engagement_rate":
                    ws_cmp.cell(row=row, column=2).number_format = '0.0"%"'
                    ws_cmp.cell(row=row, column=3).number_format = '0.0"%"'
                row += 1
        ws_cmp.column_dimensions["A"].width = 30
        ws_cmp.column_dimensions["B"].width = 18
        ws_cmp.column_dimensions["C"].width = 18
        ws_cmp.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
